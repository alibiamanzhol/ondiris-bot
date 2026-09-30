import io

from conftest import gen_bins
from ondiris_bot.bins import is_valid_bin, parse_document, parse_text

REAL = ["181240006529", "971240001315", "940140000385", "830115300573", "210440021733"]


def test_real_bins_are_valid():
    assert all(is_valid_bin(b) for b in REAL)
    assert not is_valid_bin("181240006528")
    assert not is_valid_bin("123456789012")


def test_single_bin():
    assert parse_text("181240006529").valid == ["181240006529"]


def test_hundred_bins_space_separated():
    bins = gen_bins(100)
    r = parse_text(" ".join(bins))
    assert r.valid == bins and r.invalid == []


def test_comma_and_semicolon_separated():
    assert parse_text(", ".join(REAL)).valid == REAL
    assert parse_text(";".join(REAL)).valid == REAL


def test_newlines_and_table():
    text = "Компания 1 | 181240006529\nКомпания 2 | 971240001315\n\n  940140000385  \r\n"
    assert parse_text(text).valid == REAL[:3]


def test_duplicates_removed_order_kept():
    r = parse_text("971240001315 181240006529 971240001315\n181240006529")
    assert r.valid == ["971240001315", "181240006529"]


def test_garbage_text_with_bins():
    text = ("Прошу добавить: ТОО «Альфа» (БИН 181240006529), тел. +77076400331, "
            "договор №12345 от 01.02.2026, сумма 1 500 000 тг; ИП Эверест-830115300573!!!")
    r = parse_text(text)
    assert r.valid == ["181240006529", "830115300573"]
    assert r.invalid == []


def test_invalid_values_reported():
    r = parse_text("181240006529 123456789012 12345678901 1234567890123 42")
    assert r.valid == ["181240006529"]
    assert r.invalid == ["123456789012", "12345678901", "1234567890123"]


def test_random_numbers_not_accepted():
    r = parse_text("20240101 3926909200 1000000.00 222929.900.000205 77076400331")
    assert r.valid == []


def _zero_bin() -> str:
    from ondiris_bot.bins import checksum_ok
    return next("05054000000" + str(d) for d in range(10) if checksum_ok("05054000000" + str(d)))


def test_leading_zero_restored():
    zero = _zero_bin()
    assert is_valid_bin(zero)
    assert parse_text(zero[1:]).valid == [zero]


def test_excel_xlsx_any_cell_any_sheet():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Наименование", "БИН", "Телефон"])
    ws.append(["Компания А", "181240006529", "+77076400331"])
    ws.append(["Компания Б", 971240001315, 87012345678])
    ws2 = wb.create_sheet("Лист2")
    ws2["F10"] = 940140000385.0
    ws2["A1"] = "мусор 123"
    ws2["B2"] = "830115300573; 181240006529"
    buf = io.BytesIO()
    wb.save(buf)
    r = parse_document(buf.getvalue(), "list.xlsx")
    assert r.valid == ["181240006529", "971240001315", "830115300573", "940140000385"]


def test_excel_numeric_cell_lost_leading_zero():
    from openpyxl import Workbook

    zero = _zero_bin()
    wb = Workbook()
    wb.active["A1"] = int(zero)
    buf = io.BytesIO()
    wb.save(buf)
    assert parse_document(buf.getvalue(), "x.xlsx").valid == [zero]


def test_excel_xls():
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("s")
    ws.write(0, 0, "БИН")
    ws.write(1, 3, 181240006529)
    ws.write(2, 1, "971240001315")
    buf = io.BytesIO()
    wb.save(buf)
    assert parse_document(buf.getvalue(), "old.xls").valid == ["181240006529", "971240001315"]


def test_csv_semicolon():
    data = "Название;БИН\nАльфа;181240006529\nБета;1.81240006529E+11\n".encode("cp1251")
    assert parse_document(data, "list.csv").valid == ["181240006529"]
