import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath

MAX_FILE_BYTES = 10 * 1024 * 1024

_SCI = re.compile(r"\b\d\.\d+[eE]\+?\d{1,2}\b")
_RUN = re.compile(r"(?<![\d+])\d+(?!\d)")


def checksum_ok(value: str) -> bool:
    d = [int(c) for c in value]
    for weights in (range(1, 12), (3, 4, 5, 6, 7, 8, 9, 10, 11, 1, 2)):
        s = sum(a * w for a, w in zip(d, weights)) % 11
        if s != 10:
            return s == d[11]
    return False


def is_valid_bin(value: str) -> bool:
    """12 цифр, месяц 01–12 в позициях 3–4 (общее для БИН и ИИН) и контрольная цифра."""
    return (
        len(value) == 12
        and value.isdigit()
        and 1 <= int(value[2:4]) <= 12
        and checksum_ok(value)
    )


@dataclass
class ParseResult:
    valid: list[str] = field(default_factory=list)
    invalid: list[str] = field(default_factory=list)

    def merge(self, other: "ParseResult") -> None:
        for b in other.valid:
            if b not in self.valid:
                self.valid.append(b)
        for b in other.invalid:
            if b not in self.invalid:
                self.invalid.append(b)

    def finalize(self) -> "ParseResult":
        self.invalid = [x for x in self.invalid if x not in self.valid]
        return self


def _classify(run: str, result: ParseResult) -> None:
    if len(run) == 12:
        (result.valid if is_valid_bin(run) else result.invalid).append(run)
    elif len(run) == 11 and is_valid_bin("0" + run):
        result.valid.append("0" + run)
    elif len(run) in (10, 11, 13):
        result.invalid.append(run)


def _expand_scientific(text: str) -> str:
    def repl(m: re.Match) -> str:
        try:
            return str(int(float(m.group(0))))
        except ValueError:
            return m.group(0)
    return _SCI.sub(repl, text)


def parse_text(text: str) -> ParseResult:
    raw = ParseResult()
    for run in _RUN.findall(_expand_scientific(text or "")):
        _classify(run, raw)
    result = ParseResult()
    result.merge(raw)
    return result.finalize()


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else ""
    if isinstance(value, int):
        return str(value)
    return str(value)


def _parse_cells(cells) -> ParseResult:
    result = ParseResult()
    for value in cells:
        text = _cell_to_text(value)
        if text:
            result.merge(parse_text(text))
    return result.finalize()


def _xlsx_cells(data: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                yield from row
    finally:
        wb.close()


def _xls_cells(data: bytes):
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            yield from sheet.row_values(r)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


class UnsupportedFile(Exception):
    pass


SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".pdf")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise UnsupportedFile("PDF защищён паролем. Снимите защиту или отправьте БИН текстом.")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        raise UnsupportedFile("В PDF нет текста — похоже, это скан или фото. "
                              "Отправьте Excel, текст или PDF, в котором текст можно выделить.")
    return text


def parse_document(data: bytes, filename: str) -> ParseResult:
    if len(data) > MAX_FILE_BYTES:
        raise UnsupportedFile("Файл слишком большой (максимум 10 МБ).")
    ext = PurePath(filename or "").suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _parse_cells(_xlsx_cells(data))
    if ext == ".xls":
        return _parse_cells(_xls_cells(data))
    if ext == ".csv":
        text = _decode(data)
        return _parse_cells(cell for row in csv.reader(io.StringIO(text), delimiter=_sniff(text)) for cell in row)
    if ext == ".txt":
        return parse_text(_decode(data))
    if ext == ".pdf":
        return parse_text(_pdf_text(data))
    raise UnsupportedFile("Поддерживаются файлы Excel (.xlsx, .xls), .csv, .txt и PDF.")


def _sniff(text: str) -> str:
    sample = text[:4096]
    return max((";", ",", "\t"), key=sample.count)
